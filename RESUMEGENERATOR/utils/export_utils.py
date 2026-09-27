"""
utils/export_utils.py
----------------------
Document Generation layer — Requirement 5 (Export) of the Resume Generator.
Produces downloadable TXT and DOCX bytes that Streamlit's
st.download_button can serve directly (no temp files on disk needed).
"""

import io
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def resume_to_txt(resume_sections: dict) -> bytes:
    lines = []
    for heading, body in resume_sections.items():
        lines.append(heading.upper())
        lines.append("-" * len(heading))
        lines.append(body.strip())
        lines.append("")
    return "\n".join(lines).encode("utf-8")


def resume_to_docx(resume_sections: dict, candidate_name: str) -> bytes:
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(candidate_name or "Candidate Resume")
    run.bold = True
    run.font.size = Pt(20)

    for heading, body in resume_sections.items():
        doc.add_heading(heading, level=2)
        for line in body.strip().split("\n"):
            if line.strip():
                p = doc.add_paragraph(line.strip())
                p.paragraph_format.space_after = Pt(4)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def interview_report_to_txt(report: dict) -> bytes:
    lines = ["INTERVIEW PERFORMANCE REPORT", "=" * 30, ""]
    lines.append(f"Questions Attempted: {report.get('questions_attempted')}")
    lines.append(f"Average Score: {report.get('average_score')}/10")
    lines.append("")
    lines.append("Strong Areas: " + ", ".join(report.get("strong_areas", [])))
    lines.append("Weak Areas: " + ", ".join(report.get("weak_areas", [])))
    lines.append("")
    lines.append("Overall Recommendation:")
    lines.append(report.get("overall_recommendation", ""))
    return "\n".join(lines).encode("utf-8")
